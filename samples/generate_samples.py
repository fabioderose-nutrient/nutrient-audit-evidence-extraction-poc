"""Generate the sample documents used by the demo.

Every company, address, tax ID and bank detail below is fictitious.

    pip install -r samples/requirements.txt
    python samples/generate_samples.py

Writes samples/soc-reports/*.docx (python-docx) and samples/invoices/*.pdf (reportlab).
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt
from reportlab.lib import colors
from reportlab.lib.enums import TA_RIGHT
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HERE = Path(__file__).resolve().parent

# --------------------------------------------------------------------------- assurance reports

REPORTS = [
    {
        "file": "SOC1-TypeII-Larkspur-Payroll.docx",
        "title": "SOC 1 TYPE II REPORT: LARKSPUR PAYROLL SERVICES INC.",
        "subtitle": (
            "Report on Larkspur Payroll Services Inc.'s Description of Its Payroll Processing System and on the "
            "Suitability of the Design and Operating Effectiveness of Its Controls Relevant to User Entities' "
            "Internal Control Over Financial Reporting"
        ),
        "organization": "Larkspur Payroll Services Inc. (Columbus, Ohio)",
        "short": "Larkspur",
        "system": "Larkspur Cloud Payroll and Tax Filing Platform",
        "auditor": "Halden & Pryce LLP",
        "period_label": "Review Period",
        "period": "January 1, 2026 to June 30, 2026",
        "period_phrase": "throughout the period January 1, 2026 to June 30, 2026",
        "standard": (
            "We conducted our examination in accordance with attestation standards established by the American "
            "Institute of Certified Public Accountants (AICPA), specifically SSAE No. 18, AT-C Section 320, "
            "Reporting on an Examination of Controls at a Service Organization Relevant to User Entities' "
            "Internal Control Over Financial Reporting."
        ),
        "type2": True,
        "qualified": None,
        "signoff": ("Halden & Pryce LLP", "Denver, Colorado", "August 14, 2026"),
        "description": (
            "The Larkspur platform calculates gross-to-net pay for client employees, withholds federal, state "
            "and local taxes, files quarterly tax returns on behalf of clients and releases net pay through ACH "
            "direct deposit batches."
        ),
        "objectives": [
            "CO-1 (Payroll Calculation Accuracy): Controls provide reasonable assurance that gross-to-net "
            "calculations follow the pay rules configured by each client.",
            "CO-2 (Tax Filing Completeness): Controls provide reasonable assurance that tax withholdings are "
            "remitted and filed completely and on time.",
            "CO-3 (Logical Access): Controls provide reasonable assurance that access to payroll data is "
            "restricted to authorized personnel.",
        ],
        "cuecs": [
            "Client entities approve every change to employee pay rates before the payroll cut-off.",
            "Client entities review the pre-payroll register and report discrepancies within one business day.",
        ],
        "controls": [
            ("LPS-01", "Pay rule changes require dual approval in the configuration console.",
             "Inspected a sample of 25 pay rule changes for evidence of dual approval.", "No exceptions noted."),
            ("LPS-02", "An automated reconciliation compares tax withheld with tax remitted after each payroll run.",
             "Re-performed the reconciliation for a sample of 40 payroll runs.", "No exceptions noted."),
            ("LPS-03", "Access rights to the production database are reviewed quarterly by the security team.",
             "Inspected the two quarterly access reviews performed in the period.", "No exceptions noted."),
        ],
    },
    {
        "file": "SOC2-TypeII-Copperline-Cloud.docx",
        "title": "SOC 2 TYPE II REPORT: COPPERLINE CLOUD HOSTING LTD.",
        "subtitle": (
            "Report on Copperline Cloud Hosting Ltd.'s Description of Its Managed Hosting System and on the "
            "Suitability of the Design and Operating Effectiveness of Its Controls Relevant to Security and "
            "Availability"
        ),
        "organization": "Copperline Cloud Hosting Ltd. (Portland, Oregon)",
        "short": "Copperline",
        "system": "Copperline Managed Kubernetes Hosting Service",
        "auditor": "Brightwater Assurance LLP",
        "period_label": "Review Period",
        "period": "April 1, 2025 to March 31, 2026",
        "period_phrase": "throughout the period April 1, 2025 to March 31, 2026",
        "standard": (
            "We conducted our examination in accordance with attestation standards established by the AICPA, "
            "AT-C Section 205, Assertion-Based Examination Engagements, using the criteria in TSP Section 100, "
            "2017 Trust Services Criteria for Security, Availability, Processing Integrity, Confidentiality, "
            "and Privacy (With Revised Points of Focus, 2022)."
        ),
        "type2": True,
        "qualified": (
            "Copperline states in its description that production changes are approved by a change advisory "
            "board before deployment. Our tests found that 4 of 30 sampled production changes were deployed "
            "without documented approval. As a result, controls were not operating effectively to meet the "
            "criterion CC8.1 during the period."
        ),
        "signoff": ("Brightwater Assurance LLP", "Seattle, Washington", "May 22, 2026"),
        "description": (
            "Copperline operates managed Kubernetes clusters for software companies. The service covers "
            "provisioning, patching, monitoring, backup and incident response for customer workloads hosted in "
            "two data centers."
        ),
        "objectives": [
            "CC6.1 (Logical Access): Access to the hosting infrastructure is restricted to authorized engineers.",
            "CC7.2 (Monitoring): Security events are monitored and escalated within defined time limits.",
            "CC8.1 (Change Management): Infrastructure changes are authorized, tested and approved before release.",
            "A1.2 (Availability): Backups are performed and restoration is tested.",
        ],
        "cuecs": [
            "User entities manage the accounts of their own staff in the self-service portal.",
            "User entities configure backup retention that meets their own recovery objectives.",
        ],
        "controls": [
            ("CCH-06", "Engineer access requires single sign-on with hardware-backed multi-factor authentication.",
             "Inspected the identity provider configuration and a sample of 25 sign-in logs.",
             "No exceptions noted."),
            ("CCH-08", "Production changes are approved by the change advisory board before deployment.",
             "Inspected a sample of 30 production changes for documented approval.",
             "Exception noted: 4 of 30 changes had no documented approval."),
            ("CCH-12", "Restoration from backup is tested every quarter.",
             "Inspected the four quarterly restoration test reports.", "No exceptions noted."),
        ],
    },
    {
        "file": "SOC1-TypeI-Tidewater-Clearing.docx",
        "title": "SOC 1 TYPE I REPORT: TIDEWATER CLEARING CORP.",
        "subtitle": (
            "Report on Tidewater Clearing Corp.'s Description of Its Trade Settlement System and on the "
            "Suitability of the Design of Its Controls"
        ),
        "organization": "Tidewater Clearing Corp. (Charlotte, North Carolina)",
        "short": "Tidewater",
        "system": "Tidewater Securities Settlement and Custody Platform",
        "auditor": "Marlow Kent Audit Partners LLP",
        "period_label": "Evaluation Date",
        "period": "As of March 31, 2026",
        "period_phrase": "as of March 31, 2026",
        "standard": (
            "We conducted our examination in accordance with the International Standard on Assurance "
            "Engagements (ISAE) 3402, Assurance Reports on Controls at a Service Organization, issued by the "
            "International Auditing and Assurance Standards Board."
        ),
        "type2": False,
        "qualified": None,
        "signoff": ("Marlow Kent Audit Partners LLP", "Boston, Massachusetts", "April 28, 2026"),
        "description": (
            "Tidewater matches trades received from broker-dealers, nets settlement obligations at the end of "
            "each day and instructs the custodian banks to move securities and cash."
        ),
        "objectives": [
            "CO-1 (Trade Matching): Controls provide reasonable assurance that only matched trades are settled.",
            "CO-2 (Settlement Instructions): Controls provide reasonable assurance that settlement instructions "
            "are complete and authorized.",
        ],
        "cuecs": [
            "Participants reconcile their daily settlement statement with their own trade records.",
        ],
        "controls": [
            ("TWC-11", "Unmatched trades are held in an exception queue and cannot be settled.",
             "Inspected the system configuration of the exception queue.", "Suitably designed."),
            ("TWC-14", "Settlement instructions above 5 million USD require a second approver.",
             "Inspected the approval workflow configuration.", "Suitably designed."),
        ],
    },
    {
        "file": "ISAE3000-TypeII-Fernhill-Carbon.docx",
        "title": "ISAE 3000 TYPE II ASSURANCE REPORT: FERNHILL CARBON DATA LTD.",
        "subtitle": (
            "Independent Reasonable Assurance Report on the Controls Over the Greenhouse Gas Emissions Data "
            "Service"
        ),
        "organization": "Fernhill Carbon Data Ltd. (Leeds, United Kingdom)",
        "short": "Fernhill",
        "system": "Fernhill Scope 1, 2 and 3 Emissions Reporting Service",
        "auditor": "Ashgrove Assurance Ltd",
        "period_label": "Review Period",
        "period": "January 1, 2025 to December 31, 2025",
        "period_phrase": "throughout the period January 1, 2025 to December 31, 2025",
        "standard": (
            "We conducted our engagement in accordance with International Standard on Assurance Engagements "
            "ISAE 3000 (Revised), Assurance Engagements Other than Audits or Reviews of Historical Financial "
            "Information, and, for the emissions data, ISAE 3410, Assurance Engagements on Greenhouse Gas "
            "Statements."
        ),
        "type2": True,
        "qualified": None,
        "signoff": ("Ashgrove Assurance Ltd", "Manchester, United Kingdom", "March 12, 2026"),
        "description": (
            "Fernhill collects energy, travel and procurement data from its clients, applies published emission "
            "factors and produces quarterly greenhouse gas statements for each client."
        ),
        "objectives": [
            "CO-1 (Data Completeness): Controls ensure that all activity data received is loaded and processed.",
            "CO-2 (Emission Factors): Controls ensure that emission factors are current and approved.",
        ],
        "cuecs": [
            "Client entities confirm that the activity data they submit is complete for each quarter.",
        ],
        "controls": [
            ("FCD-02", "Record counts are reconciled between the client upload and the processing database.",
             "Re-performed the reconciliation for a sample of 20 uploads.", "No exceptions noted."),
            ("FCD-05", "Emission factor tables are updated by an analyst and approved by the methodology lead.",
             "Inspected the approval records of the three factor updates in the period.", "No exceptions noted."),
        ],
    },
]


def soc_report(spec: dict, path: Path) -> None:
    doc = Document()
    doc.core_properties.author = "Sample generator"
    doc.core_properties.title = spec["title"]
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(10.5)

    doc.add_heading(spec["title"], level=0)
    doc.add_paragraph(spec["subtitle"]).runs[0].italic = True
    for label, value in (
        ("Service Organization", spec["organization"]),
        ("System Assessed", spec["system"]),
        ("Independent Service Auditor", spec["auditor"]),
        (spec["period_label"], spec["period"]),
    ):
        p = doc.add_paragraph()
        p.add_run(f"{label}: ").bold = True
        p.add_run(value)

    org = spec["organization"].split(" (")[0]
    doc.add_heading("SECTION I: INDEPENDENT SERVICE AUDITOR'S REPORT", level=1)
    doc.add_paragraph(f"To the Board of Directors of {org}:")
    doc.add_heading("Scope", level=2)
    tested = " and the operating effectiveness of those controls" if spec["type2"] else ""
    doc.add_paragraph(
        f"We have examined {org}'s (\"{spec['short']}\") description of its {spec['system']} "
        f"{spec['period_phrase']} (the description), and the suitability of the design{tested} to achieve "
        "the related control objectives stated in the description."
    )
    doc.add_heading("Service Auditor's Responsibilities", level=2)
    doc.add_paragraph(spec["standard"])
    if spec["qualified"]:
        doc.add_heading("Basis for Qualified Opinion", level=2)
        doc.add_paragraph(spec["qualified"])
    doc.add_heading("Opinion", level=2)
    if spec["qualified"]:
        doc.add_paragraph(
            "In our opinion, except for the matter described in the Basis for Qualified Opinion paragraph, "
            "in all material respects:"
        )
    else:
        doc.add_paragraph("In our opinion, in all material respects:")
    doc.add_paragraph(
        f"The description fairly presents the {spec['system']} as designed and implemented {spec['period_phrase']}.",
        style="List Bullet",
    )
    if spec["type2"]:
        doc.add_paragraph(
            "The controls stated in the description were suitably designed and operated effectively to provide "
            f"reasonable assurance that the control objectives would be achieved {spec['period_phrase']}.",
            style="List Bullet",
        )
    else:
        doc.add_paragraph(
            "The controls stated in the description were suitably designed to provide reasonable assurance that "
            f"the control objectives would be achieved if the controls operated effectively {spec['period_phrase']}.",
            style="List Bullet",
        )
    firm, city, date = spec["signoff"]
    sign = doc.add_paragraph()
    sign.add_run(firm).bold = True
    sign.add_run(f"\n{city}\n{date}")

    doc.add_heading("SECTION II: MANAGEMENT'S ASSERTION", level=1)
    doc.add_paragraph(
        f"{org} asserts that the accompanying description fairly presents the {spec['system']} "
        f"{spec['period_phrase']} and that the controls stated in the description were suitably designed"
        + (" and operated effectively." if spec["type2"] else ".")
    )
    doc.add_paragraph(f"Signed on behalf of {org}: Chief Executive Officer; Chief Financial Officer")

    doc.add_heading("SECTION III: SYSTEM DESCRIPTION", level=1)
    doc.add_paragraph(spec["description"])
    doc.add_paragraph("Control objectives evaluated:").runs[0].bold = True
    for item in spec["objectives"]:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_paragraph("Complementary user entity controls:").runs[0].bold = True
    for item in spec["cuecs"]:
        doc.add_paragraph(item, style="List Bullet")

    title = "CONTROLS, TESTS AND RESULTS" if spec["type2"] else "CONTROL DESIGN EVALUATION"
    doc.add_heading(f"SECTION IV: {title}", level=1)
    header = ("Control", "Control description", "Test performed", "Result")
    table = doc.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
    for row in spec["controls"]:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = "Sample document generated for a demo. All names are fictitious."
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.save(path)


# --------------------------------------------------------------------------- invoices

BILL_TO = [
    "Harborview Manufacturing Group Inc.",
    "400 Lakeshore Drive, Suite 1200",
    "Chicago, IL 60611, USA",
    "Tax ID / VAT: US-41-5550123",
]

INVOICES = [
    {
        "file": "northgate-freight-NGF-2026-1187.pdf",
        "supplier": "Northgate Freight Logistics Ltd.",
        "supplier_lines": ["Unit 7, Dockside Park, Felixstowe IP11 3XY, United Kingdom",
                           "accounts@northgate-freight.example | VAT ID: GB 284 1937 55"],
        "kind": "FREIGHT INVOICE",
        "meta": [("Invoice Number", "NGF-2026-1187"), ("Invoice Date", "September 18, 2026"),
                 ("Payment Terms", "Net 30"), ("Purchase Order", "PO-55012"),
                 ("Shipment", "Felixstowe to Chicago, 2 x 40ft containers")],
        "columns": ("Line", "Description", "Qty", "Unit price (GBP)", "Amount (GBP)"),
        "lines": [("010", "Ocean freight, 40ft container", "2", "2,350.00", "4,700.00"),
                  ("020", "Port handling and customs clearance", "2", "310.00", "620.00"),
                  ("030", "Inland haulage to consignee", "1", "880.00", "880.00")],
        "totals": [("Subtotal", "GBP 6,200.00"), ("VAT 0% (export of services)", "GBP 0.00"),
                   ("TOTAL DUE", "GBP 6,200.00")],
        "remittance": ["Bank transfer to: Westbury Commercial Bank",
                       "IBAN: GB82 WEST 1234 5698 7654 32", "BIC: WESTGB2L"],
        "note": None,
    },
    {
        "file": "sentinel-response-nonpo-SRT-04417.pdf",
        "supplier": "Sentinel Response Technologies Inc.",
        "supplier_lines": ["2200 Market Street, Floor 9, Philadelphia, PA 19103, USA",
                           "billing@sentinel-response.example | Tax ID: US-23-7781406"],
        "kind": "INCIDENT RESPONSE INVOICE",
        "meta": [("Invoice Number", "SRT-04417"), ("Invoice Date", "October 2, 2026"),
                 ("Payment Terms", "Due upon receipt"),
                 ("Purchase Order Reference", "NONE (non-PO emergency engagement)"),
                 ("Case Reference", "IR-2026-118")],
        "columns": ("Line", "Service", "Hours", "Rate (USD)", "Amount (USD)"),
        "lines": [("010", "Lead incident responder, triage and containment", "28.0", "240.00", "6,720.00"),
                  ("020", "Forensic analyst, log review", "22.0", "165.00", "3,630.00"),
                  ("030", "Weekend call-out fee", "1.0", "1,150.00", "1,150.00")],
        "totals": [("Subtotal", "$11,500.00"), ("Tax (professional services exempt)", "$0.00"),
                   ("NET AMOUNT PAYABLE", "$11,500.00 USD")],
        "remittance": ["Wire transfer: Granite Valley Bank | A/C: 7730418826 | Routing: 072414395"],
        "note": ("Exception note: non-PO emergency spend. A retrospective purchase order and approval by the "
                 "head of information security are required before payment."),
    },
    {
        "file": "ridgeway-robotics-RRC-30981.pdf",
        "supplier": "Ridgeway Robotics Components Corp.",
        "supplier_lines": ["88 Foundry Road, Pittsburgh, PA 15222, USA",
                           "ar@ridgeway-robotics.example | Tax ID: US-25-6604917"],
        "kind": "INVOICE",
        "meta": [("Invoice Number", "RRC-30981"), ("Invoice Date", "September 29, 2026"),
                 ("Due Date", "October 29, 2026"), ("Customer PO", "PO-55190 (spare parts, line 4)"),
                 ("Delivery Note", "DN-30981-A")],
        "columns": ("Line", "Part / description", "Qty", "Unit price (USD)", "Amount (USD)"),
        "lines": [("1", "RX-220 servo actuator", "6", "1,140.00", "6,840.00"),
                  ("2", "Harmonic gear set, 50:1", "4", "925.00", "3,700.00"),
                  ("3", "Safety-rated encoder cable, 5 m", "12", "68.50", "822.00")],
        "totals": [("Subtotal", "$11,362.00"), ("Sales tax 7.00%", "$795.34"),
                   ("Total due incl. tax", "$12,157.34")],
        "remittance": ["ACH payment: Three Rivers Savings Bank", "Account number: 5009 1182 4471",
                       "ABA routing: 043318092"],
        "note": None,
    },
    {
        "file": "kestrel-facilities-KFS-77310.pdf",
        "supplier": "Kestrel Facility Services LLC",
        "supplier_lines": ["615 West Randolph Street, Chicago, IL 60661, USA",
                           "contracts@kestrel-fs.example | Tax ID: US-36-4471925"],
        "kind": "MILESTONE INVOICE",
        "meta": [("Invoice Number", "KFS-77310"), ("Billing Date", "September 15, 2026"),
                 ("Due Date", "October 15, 2026 (Net 30)"), ("Contract PO", "PO-55233"),
                 ("Milestone", "Phase 2 of 3, HVAC replacement")],
        "columns": ("Line", "Milestone / description", "Scheduled", "% claimed", "Amount (USD)"),
        "lines": [("010", "Phase 2: rooftop unit delivery and rigging", "$42,000.00", "100%", "$42,000.00"),
                  ("020", "Phase 2: duct and controls hookup", "$18,000.00", "100%", "$18,000.00"),
                  ("030", "Retainage withheld (10%)", "-", "-10%", "-$6,000.00"),
                  ("040", "City mechanical permit, at cost", "At cost", "100%", "$250.00")],
        "totals": [("Gross progress subtotal", "$60,250.00"), ("Retainage held", "-$6,000.00"),
                   ("Sales tax (exempt, cert. EX-2219)", "$0.00"), ("NET AMOUNT REQUESTED", "$54,250.00 USD")],
        "remittance": ["Wire remittance:", "Bank: Lakefront National Bank", "A/C: 3318 2290 6604",
                       "Routing: 071928374"],
        "note": None,
    },
]


def invoice(spec: dict, path: Path) -> None:
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["Normal"], fontSize=8.5, leading=11)
    right = ParagraphStyle("right", parent=styles["Title"], alignment=TA_RIGHT, fontSize=15)
    doc = SimpleDocTemplate(str(path), pagesize=LETTER, leftMargin=18 * mm, rightMargin=18 * mm,
                            topMargin=16 * mm, bottomMargin=16 * mm, title=spec["file"],
                            author="Sample generator", invariant=1)
    story = []
    head = Table(
        [[Paragraph(f"<b>{spec['supplier']}</b><br/>" + "<br/>".join(spec["supplier_lines"]), small),
          Paragraph(spec["kind"], right)]],
        colWidths=[105 * mm, 70 * mm],
    )
    head.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [head, Spacer(1, 8 * mm)]

    meta = "<br/>".join(f"{k}: {v}" for k, v in spec["meta"])
    parties = Table(
        [[Paragraph("<b>BILL TO:</b><br/>" + "<br/>".join(BILL_TO), small),
          Paragraph("<b>INVOICE DETAILS:</b><br/>" + meta, small)]],
        colWidths=[85 * mm, 90 * mm],
    )
    parties.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story += [parties, Spacer(1, 8 * mm)]

    rows = [[Paragraph(f"<b>{c}</b>", small) for c in spec["columns"]]]
    rows += [[Paragraph(str(c), small) for c in line] for line in spec["lines"]]
    lines = Table(rows, colWidths=[14 * mm, 79 * mm, 24 * mm, 29 * mm, 29 * mm], repeatRows=1)
    lines.setStyle(TableStyle([
        ("LINEBELOW", (0, 0), (-1, 0), 0.8, colors.black),
        ("LINEBELOW", (0, -1), (-1, -1), 0.4, colors.grey),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story += [lines, Spacer(1, 6 * mm)]

    totals = "<br/>".join(f"{k}: <b>{v}</b>" for k, v in spec["totals"])
    foot = Table(
        [[Paragraph("<b>REMITTANCE:</b><br/>" + "<br/>".join(spec["remittance"]), small),
          Paragraph(totals, ParagraphStyle("tot", parent=small, alignment=TA_RIGHT))]],
        colWidths=[95 * mm, 80 * mm],
    )
    foot.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP")]))
    story.append(foot)
    if spec["note"]:
        story += [Spacer(1, 8 * mm), Paragraph(spec["note"], small)]
    story += [Spacer(1, 10 * mm),
              Paragraph("Sample document generated for a demo. All names and numbers are fictitious.", small)]
    doc.build(story)


def main() -> None:
    reports = HERE / "soc-reports"
    invoices = HERE / "invoices"
    reports.mkdir(exist_ok=True)
    invoices.mkdir(exist_ok=True)
    for spec in REPORTS:
        soc_report(spec, reports / spec["file"])
        print("wrote", reports.name + "/" + spec["file"])
    for spec in INVOICES:
        invoice(spec, invoices / spec["file"])
        print("wrote", invoices.name + "/" + spec["file"])


if __name__ == "__main__":
    main()
